from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("杨霞求职分析_交付文件")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(90, 90, 90)
LIGHT_FILL = "E8EEF5"
GRAY_FILL = "F2F4F7"
FONT_EA = "Microsoft YaHei"
FONT_LATIN = "Calibri"


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = FONT_LATIN
    run.font.size = Pt(10)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EA)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))


def set_font(run, size=None, bold=None, color=None):
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EA)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def style_doc(doc, preset="compact"):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EA)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25 if preset == "compact" else 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18 if preset == "compact" else 16, 10 if preset == "compact" else 8),
        ("Heading 2", 13, BLUE, 14 if preset == "compact" else 12, 7 if preset == "compact" else 6),
        ("Heading 3", 12, DARK_BLUE, 10 if preset == "compact" else 8, 5 if preset == "compact" else 4),
    ]:
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EA)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for list_style_name in ["List Bullet", "List Number"]:
        style = doc.styles[list_style_name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(11)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_EA)
        style.paragraph_format.left_indent = Inches(0.375 if preset == "compact" else 0.5)
        style.paragraph_format.first_line_indent = Inches(-0.188 if preset == "compact" else -0.25)
        style.paragraph_format.space_after = Pt(4 if preset == "compact" else 8)
        style.paragraph_format.line_spacing = 1.25 if preset == "compact" else 1.167


def add_title(doc, title, subtitle=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_font(run, size=18, bold=True, color=DARK_BLUE)
    if subtitle:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(10)
        run = p.add_run(subtitle)
        set_font(run, size=10.5, color=MUTED)


def add_contact_header(doc, name, target):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(name)
    set_font(run, size=18, bold=True, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(target)
    set_font(run, size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("广州 | 15627861398（微信同号） | 1427899301@qq.com")
    set_font(run, size=10, color=MUTED)


def para(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(5)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_font(r)
    else:
        r = p.add_run(text)
        set_font(r)
    return p


def bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_font(r)


def two_col_table(doc, rows):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [2700, 6660])
    table.style = "Table Grid"
    for i, (label, detail) in enumerate(rows):
        if i > 0:
            table.add_row()
        set_cell_text(table.rows[i].cells[0], label, bold=True)
        set_cell_text(table.rows[i].cells[1], detail)
        shade_cell(table.rows[i].cells[0], GRAY_FILL)
    return table


def strategy_doc():
    doc = Document()
    style_doc(doc, "brief")
    add_title(doc, "杨霞职业方向与投递策略建议", "基于既有简历、两次职业访谈与最新基础信息整理")

    doc.add_heading("一、核心判断", level=1)
    bullets(doc, [
        "她不是“只有教培经验”，而是在教培机构这个场景里长期承担了活动项目经理、课程产品经理、客户运营、内容创作者和一线交付负责人几类角色。",
        "转型时不要先追求“完美匹配岗位”。更有效的办法是用 2-3 个方向并行试投，用真实 JD 和面试反馈校准市场。",
        "展厅讲解/企业文化传播可以保留为机会型岗位，但不建议单独做主线简历，因为她本人兴趣不足，长期满意度风险较高。",
    ])

    doc.add_heading("二、建议保留的三条主线", level=1)
    rows = [
        ("主线 A：活动策划与项目执行 / 品牌活动运营", "主投方向。她有最强证据链：100+场活动、常规汇演/节庆/赛事/品牌展示、全流程统筹、现场零重大失误、沟通协调和主持表达可叠加。优先找活动运营、项目执行、品牌活动、会员活动、社群活动、门店/商场活动等岗位。"),
        ("主线 B：培训交付 / 客户培训 / 用户赋能", "次主投方向。不要只投传统 HR 培训岗，也要看产品培训、客户赋能、连锁门店培训、课程开发、讲师运营。她的课程研发、500+课件、分龄课程体系、教学交付和客户关系都能迁移。"),
        ("主线 C：内容策划 / 出镜传播 / 品牌新媒体", "机会型方向。适合投需要“会策划、能出镜、懂短视频、能用 AI 提效”的岗位；不建议投纯剪辑、低价小编或全天后期制作岗。她的价值是内容判断、表达、脚本与账号商业化，而不只是剪视频。"),
    ]
    two_col_table(doc, rows)

    doc.add_heading("三、简历组合建议", level=1)
    bullets(doc, [
        "做 3 份简历即可，不建议再做“展厅讲解专版”。展厅类岗位如果遇到特别合适的，可以用培训版或内容传播版微调标题投递。",
        "三份简历共享同一套事实，但标题、个人概述、核心能力排序和项目案例侧重点不同。",
        "所有版本都应降低“少儿口才老师”的视觉占比，把语言改成“活动项目、客户关系、培训交付、内容生产、商业化变现”。",
    ])

    doc.add_heading("四、数据口径与表达原则", level=1)
    bullets(doc, [
        "活动数据建议写成：累计策划执行 100+ 场活动，其中包含季度汇报、节庆主题、大型语言艺术赛事和品牌展示；代表性活动单场 200-500 人。这样兼容“常规活动”和“大型活动”。",
        "内容数据建议以最新信息为准：原创视频 200+ 条，全平台粉丝 3W+，小红书 2W+，线上店铺/资料产品累计收益 30 万+。",
        "客户数据建议写成：服务线下 500+ 家长/学员家庭，线上触达或服务 10000+ 客户。避免泛写 1000+ 客户导致口径混乱。",
    ])

    doc.add_heading("五、投递打法", level=1)
    bullets(doc, [
        "前 2 周用 A/B/C 三个版本并行投递：A 版 50%，B 版 30%，C 版 20%。记录投递岗位、JD 关键词、是否被查看、是否邀约、面试反馈。",
        "活动版关键词：活动运营、活动策划、项目执行、现场统筹、品牌活动、会员活动、商场活动、用户活动、亲子活动、发布会执行。",
        "培训版关键词：培训讲师、课程开发、客户培训、用户赋能、产品培训、讲师运营、培训项目执行、门店培训。",
        "内容版关键词：内容策划、短视频编导、出镜主持、品牌传播、新媒体运营、账号运营、AIGC 内容、视频号/小红书运营。",
        "避免关键词：纯销售、纯客服、纯后期剪辑、低价新媒体小编、长期外勤地推、只做讲解接待且成长空间弱的岗位。",
    ])

    doc.add_heading("六、面试叙事", level=1)
    bullets(doc, [
        "离开教培行业的说法：过去十年积累了活动统筹、培训交付、客户关系和内容传播能力，现在希望把这些能力放到更开放的商业场景里发挥。",
        "面对“没有行业经验”的说法：行业知识可以快速补，底层能力是长期训练出来的；她已经证明过自己能从 0 搭体系、做交付、做客户信任和做商业化。",
        "面对“创业经历太杂”的说法：创业恰好训练了端到端负责，不只是执行单点任务，而是能把目标拆成方案、资源、流程、现场、复盘和结果。",
    ])

    doc.add_heading("七、下一步需要补齐的素材", level=1)
    bullets(doc, [
        "整理 8-12 张活动现场照片，按“规模、舞台、观众、后台统筹、主持/讲解”分类。",
        "整理 3-5 个代表视频链接或二维码：主持、出镜讲解、短视频内容、AI 多媒体作品。",
        "把自媒体后台数据、线上店铺收益截图、家长好评截图做成作品集，脱敏后作为面试附件。",
    ])

    return doc


def common_experience(doc, focus):
    doc.add_heading("工作经历", level=1)
    para(doc, "广州棒棒星文化传媒有限公司 | 公司合伙人 / 教学主管 | 2018.06-至今")
    if focus == "event":
        bullets(doc, [
            "参与机构从 0 到成熟运营的全流程搭建，统筹运营、招生、客户维护、课程交付、团队协作和活动项目执行。",
            "作为活动总策划/总执行，累计策划执行 100+ 场活动，覆盖季度汇报演出、节庆主题活动、大型语言艺术赛事、品牌展示活动等；代表性活动单场参与 200-500 人。",
            "负责活动创意、方案撰写、场地与供应商对接、流程排期、节目编排、舞美灯光、物料、直播推广、现场主持与复盘，活动执行零重大失误。",
            "通过活动展示与口碑运营提升机构品牌影响力，早期完成 100+ 体验学员引流、50+ 正式转化，一年内学员规模破百，峰值保持 200+ 学员。",
            "运营公众号、视频号、小红书等渠道，独立完成宣传内容策划、拍摄、剪辑和发布，为活动报名、品牌露出和客户粘性提供支持。",
        ])
    elif focus == "training":
        bullets(doc, [
            "围绕 3-12 岁儿童语言表达能力，独立研发 7 个年龄分层的阶梯式课程体系，完成课程框架、教案、教材、课件和课堂活动设计。",
            "累计制作课件 500+ 份，并将音乐、多媒体投影、舞台表达和互动游戏融入课堂，提升课程趣味性、留存率和家长认可度。",
            "长期承担教学交付、课程优化、教师协作与家长沟通工作，服务线下 500+ 家长/学员家庭，建立高信任度客户关系。",
            "主导多场汇报演出、赛事展演和课堂成果展示，把培训成果转化为可感知的舞台/视频作品，增强客户续费与转介绍意愿。",
            "自学并应用短视频、音频剪辑和 AI 工具，开发线上学习资料并搭建多平台店铺，累计商业化收益 30 万+。",
        ])
    else:
        bullets(doc, [
            "负责机构多平台内容与品牌展示，覆盖公众号、视频号、小红书、抖音、B 站等渠道，独立完成选题、脚本、拍摄、出镜、剪辑、封面和发布。",
            "累计原创发布视频 200+ 条，全平台粉丝 3W+，其中小红书单平台粉丝 2W+，围绕家长、教师和语言表演爱好者需求持续输出内容。",
            "自主研发配套电子学习资料，搭建多平台线上店铺，实现内容产品商业化，累计收益 30 万+。",
            "具备主持、镜头表达、语言组织和现场应变能力，可将复杂信息转化为用户容易理解、愿意观看和愿意购买的内容。",
            "将线下舞台节目创编经验迁移到线上内容制作，综合把控文本、画面、音乐、服饰、动作、节奏和情绪表达。",
        ])

    para(doc, "新年集团 | 总经理助理 | 2017.10-2018.05")
    bullets(doc, [
        "负责部门间工作对接、会议纪要、资料整理和客户需求沟通，建立了早期职场协作、信息整理和客户沟通基础。",
    ])


def education_skills(doc, focus):
    for _ in range(4):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_before = Pt(0)
        spacer.paragraph_format.space_after = Pt(0)
    doc.add_heading("教育背景", level=1)
    para(doc, "暨南大学 新闻与传播学院 | 播音主持艺术专业 | 本科 | 2013.09-2017.06")
    bullets(doc, [
        "普通话一级乙等；广播电视播音员主持人上岗资格证 A 级。",
        "校考湖南省第一、全国第三名录取；曾获齐越朗诵艺术节三等奖及最佳选材奖。",
        "广东广播电视台新闻频道实习播音员；内地大学生优秀朗诵作品港澳展演会代表。",
    ])

    doc.add_heading("技能与工具", level=1)
    if focus == "event":
        bullets(doc, [
            "活动统筹：方案撰写、流程设计、场地/供应商对接、现场执行、主持控场、复盘总结。",
            "内容工具：剪映、Audition、WPS、图文编辑工具、多平台 AI 工具。",
            "表达能力：主持、讲解、镜头表达、客户沟通、儿童/家长/成人多对象沟通。",
        ])
    elif focus == "training":
        bullets(doc, [
            "培训能力：课程设计、教材开发、课件制作、教学交付、课堂活动设计、效果反馈。",
            "内容工具：剪映、Audition、WPS、多平台 AI 工具，可支持线上课程与培训物料制作。",
            "沟通能力：客户关系维护、家长沟通、团队协作、活动主持与培训表达。",
        ])
    else:
        bullets(doc, [
            "内容能力：选题策划、脚本撰写、出镜表达、拍摄剪辑、封面设计、账号运营、内容商业化。",
            "工具能力：剪映、Audition、图文编辑工具、WPS、多平台 AI 工具。",
            "表达能力：播音主持、活动主持、镜头表现、产品/课程讲解、用户需求洞察。",
        ])


def event_resume():
    doc = Document()
    style_doc(doc, "compact")
    add_contact_header(doc, "杨霞", "活动策划与项目执行 / 品牌活动运营")
    doc.add_heading("个人概述", level=1)
    para(doc, "近 10 年文化教育与线下活动运营经验，曾参与机构从 0 到成熟运营的全流程搭建。擅长活动策划、现场统筹、跨方沟通、主持表达和内容传播，可独立负责从创意方案、资源协调、现场执行到复盘传播的完整闭环。累计策划执行 100+ 场活动，其中代表性活动单场 200-500 人，执行稳定、细节把控强。")
    doc.add_heading("核心能力", level=1)
    bullets(doc, [
        "活动全链路执行：可独立完成活动方案、流程排期、场地/供应商对接、物料统筹、现场执行和复盘。",
        "现场统筹与应急：长期负责大型汇演、节庆活动、赛事展演等复杂场景，活动零重大失误。",
        "主持与对外表达：播音主持科班背景，普通话一级乙等，具备主持、讲解、镜头表达和现场控场能力。",
        "客户沟通与品牌传播：服务线下 500+ 家长/学员家庭，能通过活动体验和内容传播增强客户信任与品牌口碑。",
    ])
    common_experience(doc, "event")
    doc.add_heading("代表项目", level=1)
    bullets(doc, [
        "“语见未来”全国青少年语商风采大赛：担任活动总策划、现场总执行及主持人，统筹赛事流程、节目呈现、现场秩序和对外展示，提升机构品牌影响力。",
        "“语言魔方”青少年语言艺术展演：负责节目编排、场地布置、现场统筹和执行落地，活动顺利完成并获得家长高度认可。",
        "常规品牌活动体系：每年组织夏季/寒假汇报演出、六一晚会、新春晚会、元宵/母亲节/中秋等节庆活动，形成稳定活动运营机制。",
    ])
    education_skills(doc, "event")
    return doc


def training_resume():
    doc = Document()
    style_doc(doc, "compact")
    add_contact_header(doc, "杨霞", "培训交付 / 课程开发 / 客户培训与用户赋能")
    doc.add_heading("个人概述", level=1)
    para(doc, "近 10 年课程研发、教学交付与客户沟通经验，独立搭建 3-12 岁儿童语言艺术分层课程体系，累计制作课件 500+ 份。擅长把复杂内容拆解成易理解、可练习、能反馈的课程与活动，兼具主持表达、视频制作、AI 多媒体应用和客户关系维护能力，可胜任培训讲师、课程开发、客户培训、产品培训与用户赋能类岗位。")
    doc.add_heading("核心能力", level=1)
    bullets(doc, [
        "课程设计与开发：独立完成分龄课程框架、教材、课件、课堂活动和成果展示设计。",
        "培训交付与表达：播音主持专业背景，台风自然，擅长把抽象内容转化为清晰、生动、可接受的表达。",
        "客户关系与反馈闭环：长期服务线下 500+ 家长/学员家庭，能处理关切、建立信任并持续优化交付体验。",
        "线上内容与培训物料：可独立制作视频、音频、图文和 AI 多媒体培训素材，支持线上课程与学习资料开发。",
    ])
    common_experience(doc, "training")
    doc.add_heading("代表成果", level=1)
    bullets(doc, [
        "课程体系：围绕 7 个年龄层开发阶梯式语言表达课程，累计制作课件 500+ 份，并持续根据课堂反馈优化内容。",
        "培训成果展示：通过汇报演出、赛事展演、节庆活动等方式把学习成果可视化，增强学员成就感和家长信任。",
        "内容产品商业化：围绕用户学习痛点研发电子学习资料，完成内容制作、店铺搭建和销售转化，累计收益 30 万+。",
    ])
    education_skills(doc, "training")
    return doc


def content_resume():
    doc = Document()
    style_doc(doc, "compact")
    add_contact_header(doc, "杨霞", "内容策划与出镜传播 / 品牌新媒体运营")
    doc.add_heading("个人概述", level=1)
    para(doc, "播音主持专业背景，具备主持、出镜表达、短视频内容策划和多平台账号运营经验。自主运营抖音、小红书、B 站、视频号等平台，累计原创发布视频 200+ 条，全平台粉丝 3W+，小红书粉丝 2W+，并通过电子学习资料和线上店铺实现商业化收益 30 万+。适合承担内容策划、品牌传播、出镜讲解、短视频编导和 AIGC 内容制作相关工作。")
    doc.add_heading("核心能力", level=1)
    bullets(doc, [
        "内容策划与用户洞察：能围绕家长、教师、语言表演爱好者的真实痛点做选题、脚本和内容产品设计。",
        "主持与出镜表达：普通话一级乙等，广播电视播音员主持人上岗资格证 A 级，具备自然亲和的镜头表现力。",
        "短视频全流程制作：可独立完成选题、脚本、实景拍摄、画面设计、配乐混音、剪辑、封面和发布。",
        "AIGC 与商业化：熟练使用多平台 AI 工具提高内容生产效率，并有真实内容产品变现经验。",
    ])
    common_experience(doc, "content")
    doc.add_heading("代表成果", level=1)
    bullets(doc, [
        "账号增长：从 0 自主运营多平台内容矩阵，全平台粉丝 3W+，其中小红书单平台粉丝 2W+。",
        "内容生产：原创发布视频 200+ 条，覆盖语言表演、亲子学习、教师教学素材等方向。",
        "商业转化：基于用户痛点研发配套电子学习资料，搭建线上店铺并实现累计收益 30 万+。",
    ])
    education_skills(doc, "content")
    return doc


def save_doc(doc, filename):
    path = OUT_DIR / filename
    doc.save(path)
    return path


def main():
    OUT_DIR.mkdir(exist_ok=True)
    files = [
        save_doc(strategy_doc(), "杨霞职业方向与投递策略建议.docx"),
        save_doc(event_resume(), "杨霞简历_活动策划与项目执行方向_新版.docx"),
        save_doc(training_resume(), "杨霞简历_培训交付与客户赋能方向_新版.docx"),
        save_doc(content_resume(), "杨霞简历_内容策划与出镜传播方向_新版.docx"),
    ]
    for path in files:
        print(path.resolve())


if __name__ == "__main__":
    main()
