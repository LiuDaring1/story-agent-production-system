from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT_DIR = Path("杨霞求职分析_视觉优化版")

INK = RGBColor(22, 45, 71)
BLUE = RGBColor(38, 102, 166)
DEEP = RGBColor(22, 58, 95)
MUTED = RGBColor(96, 105, 118)
WHITE = RGBColor(255, 255, 255)
LIGHT_BLUE = "EAF2FB"
PALE_BLUE = "F4F8FC"
LIGHT_GRAY = "F5F7FA"
MID_BLUE = "D8E8F7"
FONT_EA = "Microsoft YaHei"
FONT_LATIN = "Calibri"
CONTENT_W = 9360


def ox(tag):
    return qn(f"w:{tag}")


def set_run(run, size=None, bold=None, color=None):
    run.font.name = FONT_LATIN
    run._element.rPr.rFonts.set(ox("ascii"), FONT_LATIN)
    run._element.rPr.rFonts.set(ox("hAnsi"), FONT_LATIN)
    run._element.rPr.rFonts.set(ox("eastAsia"), FONT_EA)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_margins(cell, top=95, start=130, bottom=95, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(ox(name))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(ox("w"), str(value))
        node.set(ox("type"), "dxa")


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(ox("fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, size=10.5, bold=False, color=None, align=None):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.12
    if align:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color or INK)


def table_geometry(table, widths, indent=0):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(ox("type"), "dxa")
    tbl_w.set(ox("w"), str(sum(widths)))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(ox("w"), str(indent))
    tbl_ind.set(ox("type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(ox("w"), str(width))
        grid.append(col)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(ox("type"), "dxa")
            tc_w.set(ox("w"), str(widths[i]))


def borders(table, color="D6DEE8", size="6"):
    tbl_pr = table._tbl.tblPr
    tbl_borders = tbl_pr.first_child_found_in("w:tblBorders")
    if tbl_borders is None:
        tbl_borders = OxmlElement("w:tblBorders")
        tbl_pr.append(tbl_borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = tbl_borders.find(ox(edge))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tbl_borders.append(tag)
        tag.set(ox("val"), "single")
        tag.set(ox("sz"), size)
        tag.set(ox("space"), "0")
        tag.set(ox("color"), color)


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(ox("eastAsia"), FONT_EA)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color, before, after in [
        ("Heading 1", 13.5, BLUE, 10, 5),
        ("Heading 2", 11.5, DEEP, 7, 3),
        ("Heading 3", 10.5, INK, 4, 2),
    ]:
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style._element.rPr.rFonts.set(ox("eastAsia"), FONT_EA)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    for name in ["List Bullet", "List Number"]:
        style = doc.styles[name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(10.2)
        style._element.rPr.rFonts.set(ox("eastAsia"), FONT_EA)
        style.paragraph_format.left_indent = Inches(0.22)
        style.paragraph_format.first_line_indent = Inches(-0.13)
        style.paragraph_format.space_after = Pt(2.5)
        style.paragraph_format.line_spacing = 1.08


def p(doc, text="", size=10.5, bold=False, color=None, align=None, after=4, before=0):
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(before)
    para.paragraph_format.space_after = Pt(after)
    para.paragraph_format.line_spacing = 1.12
    if align:
        para.alignment = align
    run = para.add_run(text)
    set_run(run, size=size, bold=bold, color=color or INK)
    return para


def add_bullets(doc, items):
    for item in items:
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(item)
        set_run(run, color=INK)


def add_section(doc, title):
    para = doc.add_heading(title, level=1)
    return para


def add_header(doc, name, target, subtitle):
    table = doc.add_table(rows=1, cols=2)
    table_geometry(table, [5750, 3610], indent=0)
    borders(table, color="FFFFFF", size="0")
    for cell in table.rows[0].cells:
        shade(cell, "163A5F")
        set_cell_margins(cell, top=165, bottom=165, start=180, end=180)
    left, right = table.rows[0].cells

    left.text = ""
    para = left.paragraphs[0]
    para.paragraph_format.space_after = Pt(2)
    run = para.add_run(name)
    set_run(run, size=21, bold=True, color=WHITE)
    para = left.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(target)
    set_run(run, size=11.5, bold=True, color=RGBColor(210, 228, 247))
    para = left.add_paragraph()
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(subtitle)
    set_run(run, size=9.6, color=RGBColor(228, 235, 242))

    right.text = ""
    for line in ["广州", "15627861398（微信同号）", "1427899301@qq.com"]:
        para = right.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        para.paragraph_format.space_after = Pt(1)
        run = para.add_run(line)
        set_run(run, size=10, color=WHITE)


def add_metric_strip(doc, metrics):
    table = doc.add_table(rows=1, cols=len(metrics))
    table_geometry(table, [CONTENT_W // len(metrics)] * len(metrics), indent=0)
    borders(table, color="C9D8E8", size="4")
    for i, (num, label) in enumerate(metrics):
        cell = table.rows[0].cells[i]
        shade(cell, LIGHT_BLUE)
        cell.text = ""
        set_cell_margins(cell, top=105, bottom=105, start=110, end=110)
        para = cell.paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(1)
        run = para.add_run(num)
        set_run(run, size=15, bold=True, color=DEEP)
        para = cell.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        run = para.add_run(label)
        set_run(run, size=8.5, bold=True, color=MUTED)


def add_keyword_bar(doc, keywords):
    table = doc.add_table(rows=1, cols=1)
    table_geometry(table, [CONTENT_W], indent=0)
    borders(table, color="D9E3EF", size="4")
    cell = table.rows[0].cells[0]
    shade(cell, PALE_BLUE)
    set_cell_margins(cell, top=90, bottom=90, start=140, end=140)
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run("岗位关键词：")
    set_run(run, size=9.5, bold=True, color=DEEP)
    run = para.add_run(" / ".join(keywords))
    set_run(run, size=9.5, color=INK)


def add_summary(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table_geometry(table, [CONTENT_W], indent=0)
    borders(table, color="D5E1EF", size="5")
    cell = table.rows[0].cells[0]
    shade(cell, LIGHT_GRAY)
    set_cell_margins(cell, top=110, bottom=110, start=150, end=150)
    cell.text = ""
    para = cell.paragraphs[0]
    para.paragraph_format.space_after = Pt(0)
    run = para.add_run(text)
    set_run(run, size=10.2, color=INK)


def add_capability_grid(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table_geometry(table, [2150, 7210], indent=0)
    borders(table, color="D6DEE8", size="4")
    for i, (label, detail) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        shade(c0, MID_BLUE)
        shade(c1, "FFFFFF")
        set_cell_text(c0, label, size=9.6, bold=True, color=DEEP)
        set_cell_text(c1, detail, size=9.6, color=INK)


def add_job(doc, title, bullets):
    para = p(doc, title, size=10.6, bold=True, color=INK, after=2, before=2)
    para.paragraph_format.keep_with_next = True
    add_bullets(doc, bullets)


def add_education(doc, focus):
    add_section(doc, "教育背景与证书")
    p(doc, "暨南大学 新闻与传播学院 | 播音主持艺术专业 | 本科 | 2013.09-2017.06", bold=True, after=2)
    add_bullets(doc, [
        "普通话一级乙等；广播电视播音员主持人上岗资格证 A 级。",
        "校考湖南省第一、全国第三名录取；齐越朗诵艺术节三等奖及最佳选材奖。",
        "广东广播电视台新闻频道实习播音员；内地大学生优秀朗诵作品港澳展演会代表。",
    ])

    add_section(doc, "工具与补充能力")
    if focus == "event":
        tools = "活动方案 / 流程排期 / 供应商与场地对接 / 现场主持 / 复盘总结 / 剪映 / Audition / WPS / AI 多媒体工具"
    elif focus == "training":
        tools = "课程设计 / 教材与课件开发 / 课堂交付 / 培训反馈 / 线上学习资料 / 剪映 / Audition / WPS / AI 多媒体工具"
    else:
        tools = "选题策划 / 脚本撰写 / 出镜表达 / 拍摄剪辑 / 封面设计 / 小红书 / 抖音 / 视频号 / B站 / AI 内容工具"
    p(doc, tools, after=2)


def add_footer_note(doc):
    p(doc, "注：以上数据基于既有工作材料、过往简历与最新补充信息整理；面试前建议结合具体岗位 JD 做 5-10 分钟关键词微调。", size=8.6, color=MUTED, after=0)


RESUMES = {
    "event": {
        "filename": "杨霞简历_活动策划与项目执行方向_视觉优化版.docx",
        "target": "活动策划与项目执行 / 品牌活动运营",
        "subtitle": "适配：活动运营、项目执行、品牌活动、用户活动、会员活动、商场/门店活动",
        "summary": "近 10 年文化教育与线下活动运营经验，曾参与机构从 0 到成熟运营的全流程搭建。核心价值不是“参与活动”，而是能把复杂活动拆成目标、流程、资源、人员、现场、传播与复盘，并稳定交付可衡量结果。",
        "metrics": [("100+场", "活动策划执行"), ("200-500人", "代表性单场规模"), ("50%+", "早期体验转化"), ("200+学员", "机构峰值规模")],
        "keywords": ["活动策划", "项目执行", "现场统筹", "流程管理", "供应商对接", "品牌活动", "用户活动", "主持控场"],
        "capabilities": [
            ("活动全流程", "从活动创意、方案撰写、流程排期、场地/供应商对接、物料统筹到现场执行、主持控场和复盘总结均可独立推进。"),
            ("复杂度与规模", "累计策划执行 100+ 场活动，覆盖汇报演出、节庆活动、赛事展演、品牌展示；代表性活动单场 200-500 人。"),
            ("结果意识", "早期通过定制化招生与体验活动完成 100+ 体验学员引流、50+ 正式转化，转化率超 50%，一年内学员规模破百。"),
            ("复合交付", "兼具主持表达、客户沟通、短视频制作和 AI 多媒体应用能力，可支持活动前宣、现场呈现和活动后传播。"),
        ],
        "jobs": [
            ("广州棒棒星文化传媒有限公司 | 公司合伙人 / 教学主管 / 活动总执行 | 2018.06-至今", [
                "主导机构从 0 到成熟运营的全流程搭建，统筹招生、课程交付、客户维护、团队协作、活动项目和品牌传播。",
                "累计策划执行 100+ 场活动，覆盖季度汇报、节庆主题、大型语言艺术赛事和品牌展示；代表性活动单场 200-500 人，现场零重大失误。",
                "负责活动方案、场地与供应商对接、流程排期、节目编排、舞美灯光、物料、直播推广、现场主持和复盘，降低活动执行风险。",
                "通过体验活动和口碑运营完成 100+ 体验学员引流、50+ 正式转化，一年内学员规模破百，峰值保持 200+ 学员。",
                "运营公众号、视频号、小红书等渠道，独立完成活动宣传内容策划、拍摄、剪辑和发布，提升活动报名与品牌露出。",
            ]),
            ("新年集团 | 总经理助理 | 2017.10-2018.05", [
                "负责部门间工作对接、会议纪要、资料整理和客户需求沟通，建立跨部门协作、信息整理和客户沟通基础。",
            ]),
        ],
        "projects": [
            ("“语见未来”全国青少年语商风采大赛", "担任活动总策划、现场总执行及主持人，统筹赛事流程、节目呈现、现场秩序与对外展示，提升机构品牌影响力。"),
            ("“语言魔方”青少年语言艺术展演", "负责节目编排、场地布置、现场统筹与执行落地，活动顺利完成并获得家长高度认可。"),
            ("常规品牌活动体系", "每年组织夏季/寒假汇报演出、六一晚会、新春晚会、元宵/母亲节/中秋等活动，形成稳定活动运营机制。"),
        ],
    },
    "training": {
        "filename": "杨霞简历_培训交付与客户赋能方向_视觉优化版.docx",
        "target": "培训交付 / 课程开发 / 客户培训与用户赋能",
        "subtitle": "适配：培训讲师、课程开发、客户培训、产品培训、用户赋能、讲师运营",
        "summary": "近 10 年课程研发、教学交付与客户沟通经验。核心价值不是“上过课”，而是能把复杂内容设计成清晰课程、可练习任务和可反馈成果，并通过沟通与成果展示持续提升客户信任。",
        "metrics": [("7个", "年龄分层课程"), ("500+份", "课件与教学物料"), ("500+家庭", "线下客户服务"), ("30万+", "内容产品收益")],
        "keywords": ["培训交付", "课程开发", "教材课件", "客户培训", "用户赋能", "效果反馈", "内容产品", "AIGC 物料"],
        "capabilities": [
            ("课程设计", "独立完成分龄课程框架、教材、课件、课堂活动与成果展示设计，能把复杂内容拆成易学、可练、可反馈的模块。"),
            ("培训交付", "播音主持专业背景，表达清晰、亲和自然，能根据儿童、家长、成人等不同对象调整表达方式。"),
            ("客户关系", "长期服务线下 500+ 家长/学员家庭，能处理关切、建立信任，并通过反馈持续优化交付体验。"),
            ("线上化能力", "可独立制作视频、音频、图文和 AI 多媒体培训素材，并有真实内容产品商业化经验。"),
        ],
        "jobs": [
            ("广州棒棒星文化传媒有限公司 | 公司合伙人 / 教学主管 / 课程负责人 | 2018.06-至今", [
                "围绕 3-12 岁儿童语言表达能力，独立研发 7 个年龄分层的阶梯式课程体系，完成课程框架、教案、教材、课件与课堂活动设计。",
                "累计制作课件 500+ 份，并将音乐、多媒体投影、舞台表达和互动游戏融入课堂，提升课程趣味性、留存率和家长认可度。",
                "长期承担教学交付、课程优化、教师协作与家长沟通，服务线下 500+ 家长/学员家庭，建立高信任度客户关系。",
                "主导汇报演出、赛事展演和课堂成果展示，把培训成果转化为可感知的舞台/视频作品，增强客户续费与转介绍意愿。",
                "自学并应用短视频、音频剪辑和 AI 工具，开发线上学习资料并搭建多平台店铺，累计商业化收益 30 万+。",
            ]),
            ("新年集团 | 总经理助理 | 2017.10-2018.05", [
                "负责部门间工作对接、会议纪要、资料整理和客户需求沟通，建立信息整理、协作推进和客户沟通基础。",
            ]),
        ],
        "projects": [
            ("分龄课程体系搭建", "围绕 7 个年龄层开发阶梯式语言表达课程，累计制作课件 500+ 份，并持续根据课堂反馈优化内容。"),
            ("培训成果展示机制", "通过汇报演出、赛事展演、节庆活动等方式把学习成果可视化，增强学员成就感和家长信任。"),
            ("线上学习资料商业化", "围绕用户学习痛点研发电子学习资料，完成内容制作、店铺搭建和销售转化，累计收益 30 万+。"),
        ],
    },
    "content": {
        "filename": "杨霞简历_内容策划与出镜传播方向_视觉优化版.docx",
        "target": "内容策划与出镜传播 / 品牌新媒体运营",
        "subtitle": "适配：内容策划、短视频编导、出镜主持、品牌传播、新媒体运营、AIGC 内容",
        "summary": "播音主持专业背景，具备主持、出镜表达、短视频内容策划和多平台账号运营经验。核心价值不是“会剪视频”，而是能围绕用户痛点做选题、脚本、表达、制作与商业化转化。",
        "metrics": [("200+条", "原创视频"), ("3W+", "全平台粉丝"), ("2W+", "小红书粉丝"), ("30万+", "内容产品收益")],
        "keywords": ["内容策划", "短视频编导", "出镜表达", "脚本撰写", "账号运营", "小红书", "视频号", "AIGC 内容"],
        "capabilities": [
            ("内容策划", "能围绕家长、教师和语言表演爱好者的真实痛点做选题、脚本和内容产品设计。"),
            ("出镜表达", "普通话一级乙等，广播电视播音员主持人上岗资格证 A 级，镜头表现亲和自然，适合讲解与品牌传播。"),
            ("制作闭环", "可独立完成选题、脚本、实景拍摄、画面设计、配乐混音、剪辑、封面和发布。"),
            ("商业转化", "能将用户需求转化为电子学习资料和线上产品，已实现累计 30 万+ 内容产品收益。"),
        ],
        "jobs": [
            ("广州棒棒星文化传媒有限公司 | 公司合伙人 / 内容负责人 / 教学主管 | 2018.06-至今", [
                "负责机构多平台内容与品牌展示，覆盖公众号、视频号、小红书、抖音、B 站等渠道，独立完成选题、脚本、拍摄、出镜、剪辑、封面和发布。",
                "累计原创发布视频 200+ 条，全平台粉丝 3W+，其中小红书单平台粉丝 2W+，围绕家长、教师和语言表演爱好者需求持续输出内容。",
                "自主研发配套电子学习资料，搭建多平台线上店铺，实现内容产品商业化，累计收益 30 万+。",
                "具备主持、镜头表达、语言组织和现场应变能力，可将复杂信息转化为用户容易理解、愿意观看和愿意购买的内容。",
                "将线下舞台节目创编经验迁移到线上内容制作，综合把控文本、画面、音乐、服饰、动作、节奏和情绪表达。",
            ]),
            ("新年集团 | 总经理助理 | 2017.10-2018.05", [
                "负责部门间工作对接、会议纪要、资料整理和客户需求沟通，建立信息整理、协作推进和客户沟通基础。",
            ]),
        ],
        "projects": [
            ("多平台账号增长", "从 0 自主运营多平台内容矩阵，全平台粉丝 3W+，其中小红书单平台粉丝 2W+。"),
            ("原创内容生产", "原创发布视频 200+ 条，覆盖语言表演、亲子学习、教师教学素材等方向。"),
            ("内容产品商业化", "基于用户痛点研发配套电子学习资料，搭建线上店铺并实现累计收益 30 万+。"),
        ],
    },
}


def build_resume(kind):
    data = RESUMES[kind]
    doc = Document()
    style_doc(doc)
    add_header(doc, "杨霞", data["target"], data["subtitle"])
    add_metric_strip(doc, data["metrics"])
    add_keyword_bar(doc, data["keywords"])
    add_section(doc, "职业概述")
    add_summary(doc, data["summary"])
    add_section(doc, "核心能力与量化证据")
    add_capability_grid(doc, data["capabilities"])
    add_section(doc, "工作经历")
    for title, items in data["jobs"]:
        add_job(doc, title, items)
    add_section(doc, "代表项目 / 成果")
    add_capability_grid(doc, data["projects"])
    add_education(doc, kind)
    add_footer_note(doc)
    return doc


def build_strategy():
    doc = Document()
    style_doc(doc)
    add_header(doc, "杨霞求职材料优化说明", "从“工作流水账”改为“量化价值证据”", "用于后续改简历、投递和面试复盘")
    add_metric_strip(doc, [("3份", "定制简历"), ("4类", "量化证据"), ("10秒", "HR首屏扫描"), ("2周", "投递复盘周期")])

    add_section(doc, "本轮优化吸收的核心原则")
    add_capability_grid(doc, [
        ("不是流水账", "简历不重点写“做了什么”，而是写“解决了什么问题、覆盖多大规模、带来什么结果”。"),
        ("量化优先", "优先呈现项目规模、用户规模、转化结果、内容产出、商业化收益、错误减少或稳定交付等可衡量证据。"),
        ("针对岗位", "三份简历分别服务活动、培训、内容三个方向；标题、关键词、首屏数字和项目案例都围绕岗位需求排序。"),
        ("AI是辅助", "AI可以帮助翻译经历、调整关键词和表达方式，但最终数据口径、事实真实性和岗位取舍必须由本人判断。"),
        ("简明专业", "视觉排版用来帮助 HR 快速抓重点，不是为了装饰；因此采用克制蓝灰色、数字条、关键词栏和能力证据矩阵。"),
    ])

    add_section(doc, "三份简历的使用建议")
    add_capability_grid(doc, [
        ("活动版", "主投。适合活动运营、项目执行、品牌活动、会员活动、商场/门店活动、亲子活动、用户活动等岗位。"),
        ("培训版", "次主投。适合培训讲师、课程开发、客户培训、用户赋能、产品培训、讲师运营和连锁门店培训。"),
        ("内容版", "机会型投递。适合内容策划、短视频编导、出镜主持、品牌传播、新媒体运营和 AIGC 内容岗位。"),
    ])

    add_section(doc, "投递前 5 分钟微调清单")
    add_bullets(doc, [
        "把目标 JD 中出现频率最高的 5-8 个关键词，替换到简历顶部“岗位关键词”栏。",
        "如果 JD 强调活动规模，就把“100+场、200-500人、零重大失误”放在最前；如果强调内容增长，就把“200+条、3W+、30万+”放在最前。",
        "如果岗位偏执行，少写宏观愿景，多写流程排期、资源协调、现场把控；如果岗位偏策划，多写选题、方案、用户洞察和复盘。",
        "面试前准备 3 个可展开案例：一个活动项目、一个客户沟通/培训交付案例、一个内容商业化案例。",
    ])
    return doc


def main():
    OUT_DIR.mkdir(exist_ok=True)
    files = [
        ("杨霞求职材料优化说明_视觉优化版.docx", build_strategy()),
        (RESUMES["event"]["filename"], build_resume("event")),
        (RESUMES["training"]["filename"], build_resume("training")),
        (RESUMES["content"]["filename"], build_resume("content")),
    ]
    for filename, doc in files:
        path = OUT_DIR / filename
        doc.save(path)
        print(path.resolve())


if __name__ == "__main__":
    main()
