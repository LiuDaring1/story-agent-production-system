from pathlib import Path

from docx import Document
from pypdf import PdfReader


OUT_DIR = Path("杨霞求职分析_工作区")

DOCX_FILES = [
    ("/Users/baiyanglin/Desktop/杨霞简历的一些基础信息.docx", "基础信息"),
    ("/Users/baiyanglin/Desktop/杨霞简历_活动策划与项目执行方向.docx", "活动策划与项目执行简历"),
    ("/Users/baiyanglin/Desktop/杨霞简历_企业培训与品牌传播方向.docx", "企业培训与品牌传播简历"),
    ("/Users/baiyanglin/Desktop/私人文件/2026年2月5日第一次谈话.docx", "2月5日谈话"),
    ("/Users/baiyanglin/Desktop/私人文件/2026年3月7日第二次谈话.docx", "3月7日谈话"),
]

PDF_FILE = ("/Users/baiyanglin/Desktop/私人文件/最诚挚的建议.pdf", "最诚挚的建议")


def extract_docx(path: str) -> str:
    doc = Document(path)
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_pdf(path: str) -> str:
    reader = PdfReader(path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def main() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    for path, label in DOCX_FILES:
        (OUT_DIR / f"{label}.txt").write_text(extract_docx(path), encoding="utf-8")
    pdf_path, pdf_label = PDF_FILE
    (OUT_DIR / f"{pdf_label}.txt").write_text(extract_pdf(pdf_path), encoding="utf-8")
    print(OUT_DIR.resolve())


if __name__ == "__main__":
    main()
