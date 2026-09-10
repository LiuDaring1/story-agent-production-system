"""Upgrade only test upstream input stubs to valid DOCX/requirements before delivery QA."""
from pathlib import Path
import json,sys,shutil
sys.path.insert(0,str(Path(__file__).resolve().parents[2]))
from docx import Document
from story_production_v2 import binding,write

def main():
    root=Path(sys.argv[1]).resolve();project=root/'fixture';status=project/'99_项目状态';cards=project/'cards';runpath=status/'story_run.json';run=json.loads(runpath.read_text())
    word=project/'manuscript.docx'
    doc=Document();doc.add_paragraph('FIXTURE');doc.add_paragraph('测试原文。');doc.save(word)
    requirements=project/'requirements.json'
    write(requirements,{'story_info':{'story_name':'FIXTURE','story_type':'测试故事','age_range':'6岁','theme_style':'离线合成夹具','sources':{k:{'kind':'offline_fixture'} for k in ('story_name','story_type','age_range','theme_style')}},'fixture_only':True})
    run['inputs']['final_word']=binding(word);run['inputs']['story_requirements']=binding(requirements);write(runpath,run)
    semantic=cards/'artifact_semantic_plan.json';payload=json.loads(semantic.read_text())
    for role in ('final_word','story_requirements'):payload['inputs'][role]=run['inputs'][role]
    write(semantic,payload)
    request=cards/'semantic_card_motion_request.json';payload=json.loads(request.read_text());payload['artifact_semantic_plan_sha256']=binding(semantic)['sha256'];write(request,payload)
    bundle=cards/'semantic_review_bundle.json';payload=json.loads(bundle.read_text());payload['artifacts']=[binding(i['path']) for i in payload['artifacts']];write(bundle,payload)
    old=cards/'independent_review.json';archive=cards/'initial_independent_review.json'
    if old.exists() and not archive.exists():shutil.copyfile(old,archive)
    review_request=cards/'independent_review_request.json';payload=json.loads(review_request.read_text());payload['artifact']=binding(bundle);payload['changed_scope']='Only final_word/story_requirements upstream input stubs replaced with valid DOCX/JSON; image/video byte hashes unchanged; review dependency rebind, not rerender';payload['previous_review']=binding(archive);write(review_request,payload)
    print(review_request)
if __name__=='__main__':main()
