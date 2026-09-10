"""Offline upstream fixtures and real R2V QA; pause for an independent review.

No production switch or mock QA. Provider artifacts are explicitly synthetic.
Usage: python tests/e2e/prepare_formal_media.py EVIDENCE_DIRECTORY
"""
from pathlib import Path
import json
import subprocess
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
from story_production_v2 import binding, write
from tests.test_formal_entry_short_media import FormalEntryShortMediaTests


def main():
    import os
    root=Path(sys.argv[1]).resolve();os.environ['STORY_E2E_EVIDENCE']=str(root)
    FormalEntryShortMediaTests.setUpClass()
    fixture=FormalEntryShortMediaTests()
    FormalEntryShortMediaTests.txt.write_text('FIXTURE\n测试原文。\n')
    FormalEntryShortMediaTests.srt.write_text('1\n00:00:00,000 --> 00:00:00,900\nFIXTURE\n\n2\n00:00:01,000 --> 00:00:02,500\n测试原文。\n')
    initial=json.loads(FormalEntryShortMediaTests.runfile.read_text())
    for role in ('confirmed_text','subtitle_txt'):
        initial['inputs'][role]=binding(FormalEntryShortMediaTests.txt)
    initial['inputs']['subtitle_srt']=binding(FormalEntryShortMediaTests.srt)
    # Other textual fixture roles originally point to the same fixture input.
    for role,item in initial['inputs'].items():
        if Path(item['path']) == FormalEntryShortMediaTests.txt:
            initial['inputs'][role]=binding(FormalEntryShortMediaTests.txt)
    FormalEntryShortMediaTests.runfile.write_text(json.dumps(initial))
    fixture.test_01_confirmed_srt_real_entry_and_drift()
    project=FormalEntryShortMediaTests.project
    status=project/'99_项目状态'; videos=project/'videos';videos.mkdir(exist_ok=True)
    for path,duration in [(videos/'S001.mp4',2),(project/'title.mp4',1)]:
        subprocess.run(['ffmpeg','-v','error','-y','-f','lavfi','-i',f'testsrc2=size=640x360:rate=24:duration={duration}', '-c:v','libx264','-crf','20','-preset','medium','-pix_fmt','yuv420p',str(path)],check=True)
    plan=project/'plan.json'
    write(plan,{'shots':[{'shot_id':'S001','source_start':1,'source_end':3,'story_text':'测试原文。','provider_duration_seconds':2}]})
    receipt=status/'mock_provider_group.json'
    write(receipt,{'schema_version':'offline-mock-provider/v1','provider':'offline-synthetic-test-fixture','production_eligible':False,
                   'shots':[{'filename':'S001.mp4','status':'downloaded','task_id':'offline-fixture-S001','output_sha256':binding(videos/'S001.mp4')['sha256']}],
                   'artifacts':[binding(videos/'S001.mp4')]})
    qa=status/'r2v_group_machine_qa.json'
    subprocess.run([sys.executable,str(Path(__file__).resolve().parents[2]/'r2v_group_qa.py'),'--plan',str(plan),'--receipt',str(receipt),'--videos-dir',str(videos),'--output',str(qa)],check=True)
    frames=[]
    for i,t in enumerate([0,.5,1,1.5,1.958]):
        frame=status/f'r2v_review_{i}.png'
        subprocess.run(['ffmpeg','-v','error','-y','-ss',str(t),'-i',str(videos/'S001.mp4'),'-frames:v','1',str(frame)],check=True)
        frames.append({**binding(frame),'time_seconds':t})
    write(status/'independent_review_request.json',{'scope':'Offline test fixture R2V, not real story visual quality; inspect actual sampled frames for moving test pattern and absence of black/frozen frames.',
        'artifact':binding(receipt),'machine_qa':binding(qa),'frames':frames,
        'review_output':str(status/'r2v_group_visual_review.json'),
        'required_review_fields':['schema_version containing independent','approved','score','critical_errors','artifact_path','artifact_sha256','reviewer_independence'],
        'not_complete':['semantic card generation and motion review','RVM preset and preview','formal assembly','dual-account rendering','directory delivery','final independent review','seal and recovery']})
    from PIL import Image, ImageDraw, ImageFont
    cards=project/'cards';cards.mkdir(exist_ok=True)
    font=ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf',48)
    card=Image.new('RGB',(640,360),(190,215,225));draw=ImageDraw.Draw(card)
    draw.text((200,70),'FIXTURE',font=font,fill=(30,50,80)); card.save(cards/'title.png')
    for i in range(24):
        frame=card.copy();d=ImageDraw.Draw(frame);x=50+i*12
        d.ellipse((x,235,x+30,265),fill=(180,100,70));frame.save(cards/f'frame_{i:03d}.png')
    subprocess.run(['ffmpeg','-v','error','-y','-framerate','24','-i',str(cards/'frame_%03d.png'),'-c:v','libx264','-crf','20','-pix_fmt','yuv420p',str(cards/'title.mp4')],check=True)
    semantic=cards/'artifact_semantic_plan.json'
    run=json.loads(FormalEntryShortMediaTests.runfile.read_text())
    from docx import Document
    manuscript=project/'manuscript.docx'
    doc=Document();doc.add_paragraph('FIXTURE');doc.add_paragraph('测试原文。');doc.save(manuscript)
    requirements=project/'requirements.json'
    write(requirements,{'story_info':{'story_name':'FIXTURE','story_type':'测试故事','age_range':'6岁','theme_style':'离线合成夹具','sources':{k:{'kind':'offline_fixture'} for k in ('story_name','story_type','age_range','theme_style')}},'fixture_only':True})
    run['inputs']['final_word']=binding(manuscript);run['inputs']['story_requirements']=binding(requirements);write(FormalEntryShortMediaTests.runfile,run)
    write(semantic,{'schema_version':'story-v2-artifact-semantics-preparation/v1','production_contract':'story-production/v2','inputs':{k:run['inputs'][k] for k in ('confirmed_text','subtitle_txt','subtitle_srt','audio','final_word','story_requirements')},'authoritative_timeline_receipt':binding(status/'timeline.json')})
    request=cards/'semantic_card_motion_request.json'
    item={'card_kind':'title_card','text':'FIXTURE','source_image_sha256':binding(cards/'title.png')['sha256'],'prompt_sha256':binding(semantic)['sha256'],'output_video_path':str(cards/'title.mp4'),'required_duration_seconds':1,'max_attempt_count':1}
    write(request,{'schema_version':'story-semantic-card-motion-request/v2','artifact_semantic_plan_sha256':binding(semantic)['sha256'],'fixture_boundary':'Mock provider artifacts, NOT ImageGen or paid generation evidence','cards':[item]})
    review_target=cards/'semantic_review_bundle.json'
    write(review_target,{'schema_version':'offline-semantic-review-bundle/v1','production_eligible':False,'artifacts':[binding(cards/'title.png'),binding(cards/'title.mp4'),binding(request),binding(semantic)],'sampled_frames':[binding(cards/f'frame_{i:03d}.png') for i in (0,12,23)],'expected_text':'FIXTURE','requirements':['Text remains exactly FIXTURE in first/middle/last frames','Only the lower dot moves smoothly horizontally; fixed camera, no tremor','No black frame or decode error; duration at least one second'],'producer_context':'offline-fixture-producer'})
    write(cards/'independent_review_request.json',{'artifact':binding(review_target),'review_output':str(cards/'independent_review.json'),'scope':'Offline semantic-card fixture only; this cannot validate ImageGen native production','required_findings':['ocr_first_frame_passed','ocr_middle_frame_passed','ocr_last_frame_passed','text_stability_passed','text_region_locked','non_text_motion_only','camera_fixed','global_jitter_passed','smooth_motion_passed','no_tremor_passed'],'no_default_approval':True})
    print(status/'independent_review_request.json')
    print(cards/'independent_review_request.json')

if __name__=='__main__':main()
