from __future__ import annotations

import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from product_package import (
    AnnotationBlock,
    DELIVERY_CJK_FONT,
    KeyingPreset,
    add_ppt_subtitle,
    clean_public_story_text,
    complete_annotation_source_lines,
    compute_source_native_layout,
    demo_crop_filter,
    fit_ppt_subtitle_text,
    load_keying_preset,
    make_blurred_background,
    preserve_native_composition,
    render_demo_preview_frame,
    render_annotation_blocks_docx,
    render_story_docx,
    reject_full_subtitle_background,
    validate_annotation_blocks,
    validate_annotation_coverage,
    validate_demo_subtitle_full_coverage,
    validate_customer_background_image,
    validate_formal_demo_logo,
)
from pptx import Presentation
from PIL import Image, ImageDraw, ImageStat
from story_project import init_project, project_paths, qa_product, write_manifest


def make_package_fixture(project: Path) -> tuple[Path, Path]:
    paths = project_paths(project)
    manifest = init_project(project, story_name="资料包QA", slug="product-qa")
    base = paths.product / "绵羊故事锦囊：资料包QA（基础版）"
    advanced = paths.product / "绵羊故事锦囊：资料包QA（进阶版）"
    base.mkdir(parents=True)
    advanced.mkdir(parents=True)
    base_names = (
        "故事文稿：资料包QA.docx",
        "故事配乐：资料包QA.mp3",
        "朗读标注：资料包QA.docx",
        "示范表演：资料包QA.mp4",
        "背景图片：资料包QA.png",
    )
    advanced_names = (
        *base_names,
        "背景视频：资料包QA（含字幕）.mp4",
        "背景视频：资料包QA（无字幕）.mp4",
        "故事PPT：资料包QA（含字幕）.pptx",
        "故事PPT：资料包QA（无字幕）.pptx",
        "A镜无人物背景视频：资料包QA.mp4",
    )
    for name in base_names:
        (base / name).write_bytes(b"fixture")
    for name in advanced_names:
        (advanced / name).write_bytes(b"fixture")
    manifest["outputs"]["product_base"] = str(base)
    manifest["outputs"]["product_advanced"] = str(advanced)
    write_manifest(paths, manifest)
    return base, advanced


class ProductQaTests(unittest.TestCase):
    def test_formal_demo_requires_official_logo_unless_reviewed_contract_disables_it(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            logo = Path(directory) / "official.png"
            logo.write_bytes(b"official-logo")
            validate_formal_demo_logo(logo, formal=True)
            validate_formal_demo_logo(None, formal=True, reviewed_no_logo=True)
            with self.assertRaisesRegex(ValueError, "官方 Logo"):
                validate_formal_demo_logo(None, formal=True)

    def test_demo_subtitles_require_opening_body_and_moral(self) -> None:
        full = ["大家好，今天讲故事。", "正文。", "这个故事告诉我们：要动脑筋。"]
        validate_demo_subtitle_full_coverage(
            ["大家好", "今天讲故事", "正文", "这个故事告诉我们", "要动脑筋"],
            full,
        )
        with self.assertRaisesRegex(ValueError, r"片头\+正文\+寓意"):
            validate_demo_subtitle_full_coverage(["正文"], full)

    def test_demo_subtitles_ignore_only_an_unspoken_standalone_document_title(self) -> None:
        validate_demo_subtitle_full_coverage(
            ["大家好，今天讲故事。", "正文。", "这个故事告诉我们：要动脑筋。"],
            ["故事标题", "大家好，今天讲故事。", "正文。", "这个故事告诉我们：要动脑筋。"],
        )

    def test_complete_annotation_source_keeps_opening_and_moral_outside_body_srt(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            story = Path(directory) / "story.txt"
            story.write_text(
                "大家好，我是绵羊姐姐。今天讲《测试》。\n"
                "正文只有这一句。\n"
                "小朋友们，这个故事告诉我们：要动脑筋。\n",
                encoding="utf-8",
            )

            self.assertEqual(
                complete_annotation_source_lines(story),
                [
                    "大家好，我是________。今天讲《测试》。",
                    "正文只有这一句。",
                    "小朋友们，这个故事告诉我们：要动脑筋。",
                ],
            )

    def test_customer_background_rejects_preblurred_and_overdense_sources_but_accepts_clear_source(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            blurred = root / "blurred.png"
            overdense = root / "overdense.png"
            clear = root / "clear.png"
            Image.new("RGB", (1280, 720), (120, 150, 180)).save(blurred)
            checker = Image.new("RGB", (1280, 720), "white")
            pixels = checker.load()
            for y in range(720):
                for x in range(1280):
                    if (x // 32 + y // 32) % 2:
                        pixels[x, y] = (20, 80, 140)
            checker.save(overdense)
            clear_image = Image.new("RGB", (1280, 720), (126, 170, 198))
            clear_pixels = clear_image.load()
            for y in range(720):
                for x in range(1280):
                    clear_pixels[x, y] = (
                        90 + int(70 * y / 719),
                        145 + int(55 * y / 719),
                        185 + int(35 * y / 719),
                    )
            # A few broad, crisp scene boundaries keep the reusable source
            # clear without making a synthetic high-frequency texture pass.
            for y in range(300, 420):
                for x in range(190, 1090):
                    if (x - 640) ** 2 + (y - 360) ** 2 < 250 ** 2:
                        clear_pixels[x, y] = (58, 125, 78)
            draw = ImageDraw.Draw(clear_image)
            for x in range(80, 1240, 145):
                draw.rounded_rectangle((x, 95, x + 42, 570), radius=18, fill=(78, 111, 72))
                draw.ellipse((x - 55, 45, x + 105, 220), fill=(72, 142, 82))
            clear_image.save(clear)
            with self.assertRaisesRegex(ValueError, "预模糊"):
                validate_customer_background_image(blurred)
            with self.assertRaisesRegex(ValueError, "高频细节过密"):
                validate_customer_background_image(overdense)
            validate_customer_background_image(clear)

    def test_background_subtitle_validation_uses_single_unmatched_srt(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            video = root / "故事背景视频（含字幕）.mp4"
            (root / "故事字幕.srt").write_text(
                "1\n00:00:00,000 --> 00:00:03,000\n大家好，我是绵羊姐姐。\n\n"
                "2\n00:00:03,000 --> 00:00:06,000\n今天我们讲一个故事。\n",
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "开头/结尾"):
                reject_full_subtitle_background(video)

    def test_source_native_layout_is_normalized_without_secondary_shrink(self) -> None:
        filter_text, x, y = compute_source_native_layout((1920, 1080), (1920, 1080), None)
        self.assertEqual(filter_text, "scale=1920:1080")
        self.assertEqual((x, y), (0, 0))
        preset = KeyingPreset(person_crop=None, person_height_ratio=0.84)
        # A full-width mode with no human-confirmed crop must still preserve
        # the source-native composition; it must not apply the old 0.84 scale.
        self.assertEqual(demo_crop_filter(Path("missing.mp4"), preset, 0.0, "full-width"), "")
        self.assertTrue(preserve_native_composition(preset, "full-width"))
        self.assertFalse(preserve_native_composition(KeyingPreset(person_crop=(0, 0, 100, 100)), "full-width"))

    def test_keying_preset_selection_is_backed_by_search_evidence(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            search = root / "keying_search.json"
            search.write_text(
                '{"candidates": [{"id": "s0.075_b0.040", "similarity": 0.075, "blend": 0.04}]}',
                encoding="utf-8",
            )
            preset = root / "keying_preset.json"
            preset.write_text(
                json.dumps(
                    {
                        "keying_candidate": "s0.115_b0.040",
                        "keying_search": str(search),
                    }
                ),
                encoding="utf-8",
            )
            with self.assertRaises(ValueError):
                load_keying_preset(preset)

    def test_background_blur_keeps_brightness_by_default(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.png"
            default = root / "default.png"
            explicit = root / "explicit.png"
            Image.new("RGB", (32, 18), (120, 150, 180)).save(source)
            make_blurred_background(source, default, 32, 18)
            make_blurred_background(source, explicit, 32, 18, brightness=0.72)
            source_mean = sum(ImageStat.Stat(Image.open(source)).mean) / 3
            default_mean = sum(ImageStat.Stat(Image.open(default)).mean) / 3
            explicit_mean = sum(ImageStat.Stat(Image.open(explicit)).mean) / 3
            self.assertAlmostEqual(default_mean, source_mean, delta=1.0)
            self.assertLess(explicit_mean, source_mean - 20)

    def test_long_ppt_subtitle_is_single_line_small_and_bounded(self) -> None:
        text = "这是一个非常非常长的故事台词，用来验证字幕在幻灯片上不会生成占满底部的粗黑条。"
        clean, font_size = fit_ppt_subtitle_text(text)
        self.assertNotIn("\n", clean)
        self.assertLessEqual(font_size, 20)
        self.assertGreaterEqual(font_size, 10)
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_ppt_subtitle(slide, text, prs)
        box = slide.shapes[-1]
        self.assertLessEqual(box.height, int(prs.slide_height * 0.09))
        self.assertLessEqual(box.top + box.height, prs.slide_height)
        self.assertEqual(box.shape_type, 13)  # PICTURE: WPS-safe transparent raster, not a black textbox.

    def test_customer_docx_uses_installed_cjk_font(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            story_docx = root / "story.docx"
            annotation_docx = root / "annotation.docx"
            render_story_docx("字体检查", "字体检查\n中文故事正文。", story_docx)
            render_annotation_blocks_docx(
                "字体检查",
                [
                    AnnotationBlock(
                        title="开场",
                        marked_text="**中文**故事正文。",
                        emotion="温暖",
                        notes=["自然地读出中文。"],
                    )
                ],
                annotation_docx,
            )
            for path in (story_docx, annotation_docx):
                with zipfile.ZipFile(path) as archive:
                    document_xml = archive.read("word/document.xml").decode("utf-8")
                self.assertIn(DELIVERY_CJK_FONT, document_xml)
                self.assertNotIn("微软雅黑", document_xml)
                self.assertNotIn("仿宋", document_xml)
            with zipfile.ZipFile(story_docx) as archive:
                story_xml = archive.read("word/document.xml").decode("utf-8")
            self.assertEqual(story_xml.count("字体检查"), 1)
            with zipfile.ZipFile(annotation_docx) as archive:
                annotation_xml = archive.read("word/document.xml").decode("utf-8")
            self.assertEqual(annotation_xml.count("w:cantSplit"), 1)
            self.assertNotIn("情绪：", annotation_xml)
            self.assertNotIn("表演细节：", annotation_xml)
            self.assertIn("温暖；自然地读出中文", annotation_xml)

    def test_demo_preview_normalizes_sought_video_timestamps(self) -> None:
        preset = KeyingPreset()
        with patch("product_package.demo_crop_filter", return_value=""), patch(
            "product_package.keying_filter_chain", return_value="[person_source]format=rgba[person_keyed]"
        ), patch(
            "product_package.preserve_native_composition", return_value=False
        ), patch("product_package.run_command") as run:
            render_demo_preview_frame(
                Path("person.mov"),
                Path("background.png"),
                Path("subtitles.mov"),
                Path("preview.png"),
                preset,
                1920,
                1080,
                0.0,
                "full-width",
                "bottom",
                37.0,
                None,
                200,
                20,
                20,
            )

        command = run.call_args.args[0]
        filters = command[command.index("-filter_complex") + 1]
        self.assertIn("[1:v]setpts=PTS-STARTPTS[person_source]", filters)
        self.assertIn("[2:v]setpts=PTS-STARTPTS,format=rgba[subtitles]", filters)

    def test_public_story_removes_host_self_introduction_entirely(self) -> None:
        self.assertEqual(clean_public_story_text("大家好，我是绵羊姐姐。"), "")
        self.assertEqual(
            clean_public_story_text("大家好，我是绵羊姐姐。今天要给大家讲故事。"),
            "今天要给大家讲故事。",
        )

    def test_annotation_requires_verbatim_text_and_natural_teacher_notes(self) -> None:
        good = [
            AnnotationBlock(
                title="决定",
                marked_text="小老虎说：“这才是**合格的评委**。”",
                emotion="认真地想明白了",
                notes=["说到“合格的评委”时放慢一点，像终于想明白了，认真地点一下头。"],
            )
        ]
        validate_annotation_blocks(good)
        validate_annotation_coverage(good, ["小老虎说：“这才是合格的评委。”"])
        with self.assertRaises(ValueError):
            validate_annotation_coverage(good, ["小老虎说：“这才是合格的小评委。”"])
        stiff = [
            AnnotationBlock(
                title="决定",
                marked_text="小老虎说：“这才是**合格的评委**。”",
                emotion="认真",
                notes=["第二个节奏落点要突出判断感，完成收束。"],
            )
        ]
        with self.assertRaises(ValueError):
            validate_annotation_blocks(stiff)

    def test_requires_both_subtitle_variants_for_video_and_ppt(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory) / "故事剪辑：资料包QA"
            _base, advanced = make_package_fixture(project)
            qa_product(project)
            report = project_paths(project).status / "qa_product_report.json"
            self.assertTrue(json.loads(report.read_text(encoding="utf-8"))["passed"])

            (advanced / "背景视频：资料包QA（无字幕）.mp4").unlink()
            qa_product(project)
            payload = json.loads(report.read_text(encoding="utf-8"))
            self.assertFalse(payload["passed"])
            self.assertTrue(any("背景视频+无字幕" in issue for issue in payload["issues"]))

    def test_product_qa_is_ledger_eligible_and_rejects_extra_customer_files(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory) / "故事剪辑：资料包QA"
            base, _advanced = make_package_fixture(project)
            qa_product(project)
            report = project_paths(project).status / "qa_product_report.json"
            payload = json.loads(report.read_text(encoding="utf-8"))
            self.assertEqual(payload["schema_version"], "story-product-machine-qa/v2")
            self.assertEqual(payload["critical_errors"], [])

            (base / "多余说明.txt").write_text("不应进入客户包", encoding="utf-8")
            qa_product(project)
            payload = json.loads(report.read_text(encoding="utf-8"))
            self.assertFalse(payload["passed"])
            self.assertTrue(payload["critical_errors"])
            self.assertTrue(any("文件数应为 5" in issue for issue in payload["issues"]))

    def test_rejects_internal_reports_inside_customer_packages(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            project = Path(directory) / "故事剪辑：报告泄漏"
            base, _advanced = make_package_fixture(project)
            (base / "成本报告.md").write_text("internal", encoding="utf-8")
            qa_product(project)
            payload = json.loads((project_paths(project).status / "qa_product_report.json").read_text(encoding="utf-8"))
            self.assertFalse(payload["passed"])
            self.assertTrue(any("混入内部报告" in issue for issue in payload["issues"]))


if __name__ == "__main__":
    unittest.main()
